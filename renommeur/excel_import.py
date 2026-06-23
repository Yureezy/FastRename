from __future__ import annotations

from dataclasses import dataclass

DEFAULT_TEXT_COLUMN = "fichier"
DEFAULT_IMAGE_COLUMN = "photo"


@dataclass
class ImportedReference:
    name: str
    image_bytes: bytes | None = None
    image_ext: str = ".png"


def _image_bytes(img) -> bytes | None:
    getter = getattr(img, "_data", None)
    try:
        if callable(getter):
            return getter()
    except Exception:  # noqa: BLE001
        pass
    ref = getattr(img, "ref", None)
    if hasattr(ref, "getvalue"):
        return ref.getvalue()
    if hasattr(ref, "read"):
        return ref.read()
    return None


def read_sheets(xlsx_path) -> dict[str, list[str]]:
    """Renvoie {nom d'onglet: [en-têtes de la ligne 1]} — lecture rapide (read_only)."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True)
    result: dict[str, list[str]] = {}
    for ws in wb.worksheets:
        headers: list[str] = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            headers = [str(c).strip() for c in row if c is not None and str(c).strip()]
            break
        result[ws.title] = headers
    wb.close()
    return result


def import_references(
    xlsx_path,
    text_column: str = DEFAULT_TEXT_COLUMN,
    image_column: str | None = DEFAULT_IMAGE_COLUMN,
    sheet_name: str | None = None,
) -> tuple[list[ImportedReference], list[str]]:
    from openpyxl import load_workbook

    warnings: list[str] = []
    wb = load_workbook(xlsx_path)

    def find_col(sheet, name: str | None) -> int | None:
        if not name:
            return None
        for cell in sheet[1]:
            if cell.value and str(cell.value).strip().casefold() == name.casefold():
                return cell.column
        return None

    # Onglet imposé par l'utilisateur, sinon celui qui contient la colonne texte.
    ws = None
    if sheet_name:
        ws = next((s for s in wb.worksheets if s.title == sheet_name), None)
    if ws is None:
        ws = wb.active
        best = -1
        for sheet in wb.worksheets:
            score = (2 if find_col(sheet, text_column) else 0) + (
                1 if find_col(sheet, image_column) else 0
            )
            if score > best:
                best, ws = score, sheet
        if best <= 0:
            ws = wb.active
            warnings.append(
                f"Aucun onglet ne contient la colonne « {text_column} » : onglet actif utilisé."
            )

    text_col = find_col(ws, text_column)
    if text_col is None:
        text_col = 1
        warnings.append(f"Colonne « {text_column} » introuvable : 1re colonne utilisée.")
    image_col = find_col(ws, image_column)

    names_by_row: dict[int, str] = {}
    order: list[int] = []
    for r in range(2, ws.max_row + 1):
        val = ws.cell(row=r, column=text_col).value
        if val is not None and str(val).strip():
            names_by_row[r] = str(val).strip()
            order.append(r)

    # Ancrage d'image : ligne 0-based -> ligne Excel = row + 1.
    # Si aucune colonne image n'est demandée, on n'extrait pas d'images.
    imgs = (getattr(ws, "_images", []) or []) if image_column else []
    images_in_row: dict[int, list[tuple[int, bytes, str]]] = {}
    for img in imgs:
        frm = getattr(getattr(img, "anchor", None), "_from", None)
        if frm is None:
            continue
        data = _image_bytes(img)
        if not data:
            continue
        ext = "." + (getattr(img, "format", None) or "png").lower()
        images_in_row.setdefault(frm.row + 1, []).append((frm.col, data, ext))

    if image_column:
        if not imgs:
            warnings.append(
                "Aucune image intégrée détectée. Si tes images sont insérées « dans » les "
                "cellules (fonction récente d'Excel), colle-les plutôt en image classique "
                "« sur » la feuille pour qu'elles soient lisibles."
            )
        elif image_col is None:
            warnings.append(
                f"Colonne « {image_column} » introuvable : l'image la plus à gauche de "
                "chaque ligne est utilisée."
            )

    def pick(row: int) -> tuple[int, bytes, str] | None:
        candidates = images_in_row.get(row)
        if not candidates:
            return None
        # Plusieurs images possibles par ligne (photo + QR) : on prend la plus
        # proche de la colonne demandée, pour ne pas confondre photo et QR code.
        if image_col is not None:
            target = image_col - 1
            candidates = sorted(candidates, key=lambda c: abs(c[0] - target))
        else:
            candidates = sorted(candidates, key=lambda c: c[0])
        return candidates[0]

    refs: list[ImportedReference] = []
    for r in order:
        img = pick(r)
        if img:
            refs.append(ImportedReference(names_by_row[r], img[1], img[2]))
        else:
            refs.append(ImportedReference(names_by_row[r]))
    return refs, warnings
