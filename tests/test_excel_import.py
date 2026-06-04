"""Tests de l'import Excel (choix de l'onglet, colonne texte, colonne image)."""

import io

import pytest

pytest.importorskip("openpyxl")
pytest.importorskip("PIL")

from openpyxl import Workbook  # noqa: E402
from openpyxl.drawing.image import Image as XLImage  # noqa: E402
from PIL import Image as PILImage  # noqa: E402

from renommeur.excel_import import import_references  # noqa: E402


def _make_xlsx(path, with_images=True):
    """Onglet unique : colonne A = Photo (images), colonne B = Fichier (texte)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Photos"
    ws["A1"] = "Photo"
    ws["B1"] = "Fichier"
    names = ["VD 1 JKT LEROY_", "VD 1 OSH 1929 K_", "VD 1 OSH KUST_"]
    for i, name in enumerate(names, start=2):
        ws.cell(row=i, column=2, value=name)
        if with_images:
            buf = io.BytesIO()
            PILImage.new("RGB", (16, 16), ["red", "green", "blue"][i - 2]).save(buf, "PNG")
            buf.seek(0)
            ws.add_image(XLImage(buf), f"A{i}")
    wb.save(path)
    return names


def test_import_texte_et_images(tmp_path):
    xlsx = tmp_path / "refs.xlsx"
    names = _make_xlsx(xlsx, with_images=True)

    refs, warnings = import_references(xlsx)

    assert [r.name for r in refs] == names
    assert all(r.image_bytes for r in refs)
    assert not warnings


def test_import_sans_image_avertit(tmp_path):
    xlsx = tmp_path / "refs.xlsx"
    names = _make_xlsx(xlsx, with_images=False)

    refs, warnings = import_references(xlsx)

    assert [r.name for r in refs] == names
    assert all(r.image_bytes is None for r in refs)
    assert warnings  # avertit qu'aucune image n'a été trouvée


def test_choisit_le_bon_onglet_et_la_colonne_photo(tmp_path):
    # Onglet actif "Info" sans la colonne ; onglet "Photos" avec Fichier + Photo + QR.
    wb = Workbook()
    info = wb.active
    info.title = "Info"
    info["A1"] = "Type de photo"

    photos = wb.create_sheet("Photos")
    photos["B1"] = "Photo"
    photos["D1"] = "QR Code"
    photos["E1"] = "Fichier"
    names = ["VD 1 JKT LEROY_", "VD 1 OSH 1929 K_"]
    for i, name in enumerate(names, start=2):
        photos.cell(row=i, column=5, value=name)  # colonne E = Fichier
        for col, size in ((2, 40), (4, 8)):  # B = photo (grande), D = QR (petite)
            buf = io.BytesIO()
            PILImage.new("RGB", (size, size), "red").save(buf, "PNG")
            buf.seek(0)
            photos.add_image(XLImage(buf), f"{chr(64 + col)}{i}")

    xlsx = tmp_path / "multi.xlsx"
    wb.save(xlsx)

    refs, warnings = import_references(xlsx)  # défauts: text="fichier", image="photo"

    assert [r.name for r in refs] == names  # bon onglet, bonne colonne texte
    # L'image retenue est la PHOTO (40x40), pas le QR code (8x8).
    for r in refs:
        assert r.image_bytes
        assert PILImage.open(io.BytesIO(r.image_bytes)).size == (40, 40)


def test_colonne_texte_introuvable_utilise_premiere(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Donnees"
    ws["A1"] = "produits"
    ws["B1"] = "autre"
    for i, name in enumerate(["AAA", "BBB"], start=2):
        ws.cell(row=i, column=1, value=name)
    xlsx = tmp_path / "x.xlsx"
    wb.save(xlsx)

    refs, warnings = import_references(xlsx, text_column="fichier", image_column="photo")

    assert [r.name for r in refs] == ["AAA", "BBB"]  # repli sur la 1re colonne
    assert any("introuvable" in w for w in warnings)
